"""Stage 21: Deep-dive report for the 'amount_deducted_twice' sub-disposition -- size of
problem, primary channel, FCR, and the full contacts-to-resolution distribution (not just
touch-count buckets from Stage 09, but real elapsed-time-to-resolution per episode).

Resolution-episode logic reused from 09_resolution_fcr.py's 48h-silence-gap rule, but rebuilt
here (rather than reading resolution_by_issue.parquet directly) because we also need each
episode's first-touch channel and its real elapsed hours (last touch - first touch), neither of
which Stage 09 keeps.

Reads:  intermediate/run_2026-05-01_2026-08-22/dedupable_raw.parquet
Writes: output/PRX_AmountDeductedTwice_DeepDive_2026-05-01_to_2026-08-22.xlsx
"""
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from pipeline_utils import PIPELINE_ROOT

TARGET = "amount_deducted_twice"
WINDOW_DIR = PIPELINE_ROOT / "intermediate" / "run_2026-05-01_2026-08-22"
WINDOW_HOURS = 48

HEADER_FILL = "632423"


def autosize_and_freeze(writer, sheet_name, df):
    ws = writer.sheets[sheet_name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        width = min(60, max(10, int(df[col].astype(str).map(len).clip(upper=60).max() if len(df) else 10) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width


def style_header(writer, sheet_name, n_cols):
    ws = writer.sheets[sheet_name]
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_episodes(order_df):
    """One order's contacts, sorted by time -> list of resolution episodes (48h-gap rule,
    same boundary logic as 09_resolution_fcr.py's contacts_to_resolution)."""
    g = order_df.sort_values("created_at")
    times = g["created_at"].tolist()
    channels = g["channel"].tolist()
    sub_channels = g["sub_channel"].tolist()
    window = pd.Timedelta(hours=WINDOW_HOURS)

    episodes = []
    ep_start_idx = 0
    for i in range(1, len(times)):
        if times[i] - times[i - 1] > window:
            episodes.append((ep_start_idx, i - 1))
            ep_start_idx = i
    episodes.append((ep_start_idx, len(times) - 1))

    rows = []
    for start, end in episodes:
        touches = end - start + 1
        elapsed_hours = (times[end] - times[start]).total_seconds() / 3600
        rows.append({
            "touches": touches,
            "elapsed_hours": elapsed_hours,
            "first_touch_channel": channels[start],
            "first_touch_sub_channel": sub_channels[start],
            "episode_start": times[start],
            "episode_end": times[end],
        })
    return rows


def main():
    dedupable = pd.read_parquet(WINDOW_DIR / "dedupable_raw.parquet")
    dedupable["created_at"] = pd.to_datetime(dedupable["created_at"])
    total_contacts_all = len(dedupable)
    print(f"[21] total contacts, all dispositions, window: {total_contacts_all}")

    sub = dedupable[dedupable["normalized_sub_disposition"] == TARGET].copy()
    sub = sub.dropna(subset=["final_order_id"])
    sub["final_order_id"] = sub["final_order_id"].astype("int64")
    total_contacts = len(sub)
    total_orders = sub["final_order_id"].nunique()
    ipi = total_contacts / total_orders
    share = total_contacts / total_contacts_all
    print(f"[21] {TARGET}: contacts={total_contacts} orders={total_orders} IPI={ipi:.3f} share={share:.4%}")

    # --- primary (first-touch) channel, at the ORDER level ---
    first_touch = sub.sort_values("created_at").groupby("final_order_id").first()
    primary_channel_counts = first_touch["channel"].value_counts()
    primary_subchannel_counts = first_touch["sub_channel"].value_counts()

    # --- resolution episodes (48h-gap rule, with elapsed time + first-touch channel) ---
    all_episodes = []
    for order_id, g in sub.groupby("final_order_id"):
        for ep in build_episodes(g):
            ep["order_id"] = order_id
            all_episodes.append(ep)
    ep_df = pd.DataFrame(all_episodes)
    n_episodes = len(ep_df)

    fcr = (ep_df["touches"] == 1).mean()
    touch_percentiles = ep_df["touches"].quantile([0.5, 0.75, 0.85, 0.95, 0.99, 1.0])
    median_resolution_hours_all = ep_df["elapsed_hours"].median()
    multi = ep_df[ep_df["touches"] > 1]
    median_resolution_hours_multi = multi["elapsed_hours"].median() if len(multi) else float("nan")

    print(f"[21] resolution episodes: {n_episodes}, FCR (1-touch): {fcr:.2%}")
    print(f"[21] touches-to-resolution percentiles (p50/p75/p85/p95/p99/p100): "
          f"{touch_percentiles.values}")
    print(f"[21] median resolution time, ALL episodes: {median_resolution_hours_all:.1f}h; "
          f"MULTI-touch episodes only: {median_resolution_hours_multi:.1f}h "
          f"({len(multi)}/{n_episodes} episodes needed 2+ touches)")

    # --- write report ---
    out_path = PIPELINE_ROOT / "output" / "PRX_AmountDeductedTwice_DeepDive_2026-05-01_to_2026-08-22.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        summary = pd.DataFrame([
            {"Metric": "Total contacts, ALL dispositions (window)", "Value": total_contacts_all},
            {"Metric": "Total contacts, amount_deducted_twice", "Value": total_contacts},
            {"Metric": "Share of all contacts", "Value": f"{share:.4%}"},
            {"Metric": "Total distinct orders", "Value": total_orders},
            {"Metric": "Interactions per Issue (IPI = contacts/orders)", "Value": round(ipi, 3)},
            {"Metric": "Primary (first-touch) channel", "Value": primary_channel_counts.idxmax()},
            {"Metric": "Primary channel share of orders", "Value": f"{primary_channel_counts.max()/total_orders:.1%}"},
            {"Metric": "Primary (first-touch) sub-channel", "Value": primary_subchannel_counts.idxmax()},
            {"Metric": "Resolution episodes (Order x 48h-gap)", "Value": n_episodes},
            {"Metric": "FCR (order level, 1-touch episodes)", "Value": f"{fcr:.2%}"},
            {"Metric": "Median contacts-to-resolution (touches)", "Value": touch_percentiles[0.5]},
            {"Metric": "p75 contacts-to-resolution", "Value": touch_percentiles[0.75]},
            {"Metric": "p85 contacts-to-resolution", "Value": touch_percentiles[0.85]},
            {"Metric": "p95 contacts-to-resolution", "Value": touch_percentiles[0.95]},
            {"Metric": "p99 contacts-to-resolution", "Value": touch_percentiles[0.99]},
            {"Metric": "p100 (max) contacts-to-resolution", "Value": touch_percentiles[1.0]},
            {"Metric": "Median resolution time (hours), ALL episodes", "Value": round(median_resolution_hours_all, 1)},
            {"Metric": "Median resolution time (hours), MULTI-touch episodes only", "Value": round(median_resolution_hours_multi, 1)},
            {"Metric": "Share of episodes needing 2+ touches", "Value": f"{len(multi)/n_episodes:.2%}"},
        ])
        summary.to_excel(writer, sheet_name="Summary", index=False)
        autosize_and_freeze(writer, "Summary", summary)
        style_header(writer, "Summary", len(summary.columns))

        chan = pd.DataFrame({
            "Channel": primary_channel_counts.index,
            "Orders (first touch)": primary_channel_counts.values,
            "Share": (primary_channel_counts / total_orders).values,
        })
        chan.to_excel(writer, sheet_name="Primary Channel", index=False)
        ws = writer.sheets["Primary Channel"]
        for r in range(2, len(chan) + 2):
            ws.cell(row=r, column=3).number_format = "0.00%"
        autosize_and_freeze(writer, "Primary Channel", chan)
        style_header(writer, "Primary Channel", len(chan.columns))

        ep_out = ep_df.sort_values("touches", ascending=False)[
            ["order_id", "touches", "elapsed_hours", "first_touch_channel",
             "first_touch_sub_channel", "episode_start", "episode_end"]
        ]
        ep_out.to_excel(writer, sheet_name="Resolution Episodes", index=False)
        autosize_and_freeze(writer, "Resolution Episodes", ep_out)
        style_header(writer, "Resolution Episodes", len(ep_out.columns))

    print(f"[21] wrote {out_path}")


if __name__ == "__main__":
    main()
