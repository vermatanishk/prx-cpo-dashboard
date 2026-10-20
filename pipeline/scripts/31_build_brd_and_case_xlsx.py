import csv, re, json, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC='/Users/tanishkverma/Desktop/CTO/Tech_Yodas_Payments_Exhaustive_Report.csv'
NOTES='/private/tmp/claude-501/-Users-tanishkverma-Desktop-CTO/e71f6b02-760f-43b0-9a91-e0cd09f5bcf4/scratchpad/notes.jsonl'
OUT='/Users/tanishkverma/Desktop/CTO/deliverables/Payment_Cases_L30_Tech_Yodas.xlsx'

rows=list(csv.DictReader(open(SRC)))
notes={}
for l in open(NOTES):
    r=json.loads(l); notes[r['tid']]=r

def scrub(s):
    s=re.sub(r'\b\d{10}\b','[phone]',s)
    s=re.sub(r'\b\d{5} \d{5}\b','[phone]',s)
    s=re.sub(r'\b[\w.\-]+@(?:ybl|oksbi|okaxis|ptyes|ibl|upi|axl|waaxis|cashfreensdlpb|mairtel|fam|[a-z]{2,10})\b','[upi]',s)
    return s

ORDER_OVERRIDE={
 'yHpZj2aIqso':'(not available - payment mis-tagged to order 3119032)',
 '7IL8cBD7cKU':'(not available - cart 6973124)',
 '9KEhdr0AW9g':'4279815',
 '_y7oKwFnNQg':'25 listed; 99 affected',
 'WARS4kf2rzg':'3733958',
}
TYPE_LABEL={
 'G1':'Amount deducted twice / duplicate debit',
 'G2':'Amount deducted, order not placed / not found / not linked',
 'G3':'Refund not credited (payout failed / stuck at gateway)',
 'G4':'Refund paid but status shown wrong',
 'G5':'Refund amount discrepancy / shortfall',
 'G6':'Return refund blocked / return-order data',
 'G7':'COD to online payment auto-refunded',
 'G8':'Payment failed / gateway incident',
 'G9':'Wallet / Edge fee paid but not provisioned',
}
def case_no(r):
    t=r['ticket']
    if t and t.startswith(('I','T')) and re.match(r'^[IT]\d+$',t): return t
    return 'No ticket ('+r['thread_id']+')'
def orders(r):
    if r['thread_id'] in ORDER_OVERRIDE: return ORDER_OVERRIDE[r['thread_id']]
    ids=re.findall(r'\b\d{7}\b',r['order_ids'])
    seen=[]; [seen.append(i) for i in ids if i not in seen]
    return ', '.join(seen) if seen else '(not available)'
NEED={'V1':'','V2':' View would be solvable (Y) once gateway-side payment/refund events are visible.',
      'T':' View can pre-diagnose and prefill the escalation; the fix itself needs tech/payments.',
      'P':' Not a payment-state question (limit / permission / product defect).'}
data=[]
for r in rows:
    n=notes[r['thread_id']]
    date=r['raised_ist'][:10]
    rem=(f"Raised {date}. Reported: {r['reported_issue']} | What went wrong: {r['what_actually_went_wrong']} | "
         f"Tech / Sherlock: {r['tech_or_sherlock_action']} | Time to close: {r['time_to_resolve']}.{NEED[r['bucket']]}")
    if r['clickhouse_payment_rows_for_order']:
        rem+=f" | Our payment table for this order: {r['clickhouse_payment_rows_for_order']}."
    data.append((case_no(r),orders(r),TYPE_LABEL[r['group'][:2]],'Y' if r['bucket']=='V1' else 'N',scrub(rem),r['bucket'],r['raised_ist']))

wb=Workbook(); ws=wb.active; ws.title='Payment cases L30'
BL={'V1':'1. Solved by view on data we already hold','V2':'2. Solved once gateway-side events are visible','T':'3. View diagnoses; fix needs Tech / Payments','P':'4. Product / policy (not a visibility problem)'}
ENG={'V1':'N','V2':'N (after gateway data added)','T':'Y','P':'Y'}
hdr=['Case #','OrderID (if available)','Issue type','Solvable by view (Y/N)','View bucket (1-4)','Engineering needed once views exist (Y/N)','Remarks and details']
ws.append(hdr)
for d in data: ws.append([d[0],d[1],d[2],d[3],BL[d[5]],ENG[d[5]],d[4]])
thin=Side(style='thin',color='BBBBBB')
for c in ws[1]:
    c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(wrap_text=True,vertical='center')
widths=[18,22,38,14,40,18,120]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    row[3].alignment=Alignment(horizontal='center',vertical='top')
    row[3].fill=PatternFill('solid',fgColor='C6EFCE' if row[3].value=='Y' else 'FCE4D6')
    row[5].alignment=Alignment(horizontal='center',vertical='top',wrap_text=True)
    row[4].fill=PatternFill('solid',fgColor={'1':'C6EFCE','2':'FFF2CC','3':'FCE4D6','4':'D9D9D9'}[row[4].value[0]])
ws.freeze_panes='A2'; ws.auto_filter.ref=f"A1:G{ws.max_row}"

# summary sheet
s=wb.create_sheet('Summary')
tot=len(data); y=sum(d[3]=='Y' for d in data)
v2=sum(d[5]=='V2' for d in data); t=sum(d[5]=='T' for d in data); p=sum(d[5]=='P' for d in data)
s.append(['Window','18 Sep - 17 Oct 2026 (Tech_yodas Google Chat space)']); s.append([])
s.append(['Measure','Cases','% of total'])
for lab,v in [('Payment cases raised in last 30 days',tot),('Solvable by view using order and payment data already on record (Y)',y),
              ('Solvable only if gateway-side payment/refund events are also shown',v2),
              ('View pre-diagnoses; fix needs tech / payments action',t),('Product / policy / permission - not addressable by a view',p)]:
    s.append([lab,v,round(v/tot,3)])
s.append([]); s.append(['View bucket','Cases','% of total','Engineering needed?'])
for k in ['V1','V2','T','P']:
    n=sum(d[5]==k for d in data); s.append([BL[k],n,round(n/tot,3),ENG[k]])
s.append(['Engineering NOT needed once views exist (1+2)',y+v2,round((y+v2)/tot,3),'N'])
s.append(['Engineering still needed (3+4)',t+p,round((t+p)/tot,3),'Y'])
s.append([]); s.append(['Issue type','Cases','Solvable by view (Y)','N'])
by=collections.defaultdict(lambda:[0,0])
for d in data:
    by[d[2]][0]+=1; by[d[2]][1]+= d[3]=='Y'
for k,(a,b) in sorted(by.items(),key=lambda x:-x[1][0]): s.append([k,a,b,a-b])
s.append([]); s.append(['Weekly volume (by raised date)','Cases'])
import datetime as dt
wk=collections.Counter(((dt.datetime.strptime(d[6][:10],'%Y-%m-%d').date()-dt.date(2026,9,18)).days)//7 for d in data)
for k,lab in enumerate(['18-24 Sep','25 Sep-1 Oct','2-8 Oct','9-15 Oct','16-17 Oct (2 days)']): s.append([lab,wk.get(k,0)])
s.column_dimensions['A'].width=70
for c in 'BCD': s.column_dimensions[c].width=22
for c in s[3]: c.font=Font(bold=True)

# excluded sheet
x=wb.create_sheet('Excluded (not payment)')
x.append(['Case #','OrderID (if available)','Why excluded','Remarks'])
for tid,n in sorted(notes.items(),key=lambda kv:kv[1]['raised']):
    if n['cat'].startswith('OUT OF SCOPE'):
        ids=', '.join(re.findall(r'\b\d{7}\b',n['orders'])) or '(not available)'
        x.append([n['ticket'] if re.match(r'^[IT]\d+$',n['ticket']) else 'No ticket ('+tid+')',ids,n['cat'].replace('OUT OF SCOPE ','').strip('()'),scrub(f"Raised {n['raised'][:10]}. {n['issue']} {n['cause']}")])
for c in x[1]: c.font=Font(bold=True)
for i,w in enumerate([18,22,40,90],1): x.column_dimensions[get_column_letter(i)].width=w
wb.save(OUT)
print(tot,y,v2,t,p, len(wb['Excluded (not payment)']['A'])-1)
