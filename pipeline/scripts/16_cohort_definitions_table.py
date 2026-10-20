"""Stage 16: Rebuild the Cohort Definitions reference table (S0-S3, M0-M12, MX) with the
current, bug-fixed counts -- same layout as the user's earlier "Cohorts - audit" sheet
(# | First contact | EDD change | Second+ contact | Count | Overall share | Root cause |
Fix | Owner), refreshed after two changes since that sheet was built:

1. M-codes are now decided by contact 1 vs contact 2 ONLY (contact 3+ out of scope), one
   code applied to every row of that order -- per the user's explicit override.
2. The "Second+ contact" (breach) check now compares contact 2 against the EDD that was
   ACTUALLY active at contact 2's own moment, not against the stale Original EDD. This
   fixed a real bug: if the EDD got revised between contact 1 and contact 2, contact 2 could
   show "after Original EDD" while genuinely landing within the newly revised promise. That
   bug had collapsed into M9 specifically; fixing it splits M9's old "After OG EDD / No
   Change (as of 1st contact) / breach" bucket into two: M9 (still breaching) and the new
   M12 (self-corrected before contact 2, no breach) -- a case that was previously impossible
   to detect and always miscounted as M9.

Reads:  /tmp/edd_with_cohort_v3_fixed.parquet (contact-level detail + per-row cohort code,
        produced by the verified Python replica of the Apps Script logic)
Writes: pipeline/output/PRX_Cohort_Definitions_Updated.xlsx
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from pipeline_utils import PIPELINE_ROOT

GREEN = "C6E0B4"
YELLOW = "FFE699"
ORANGE = "F8CBAD"
GREY = "D9D9D9"
HEADER_FILL = "632423"

ROWS = [
    # code, first_contact, edd_change, second_contact, root_cause, fix, owner
    ("S1", "Before OG EDD", "No Change", "No Second contact",
     "Pure anxiety-driven contact",
     "Live, trustworthy tracker + proactive WhatsApp on status", "Product"),
    ("S2", "On OG EDD", "No Change", "No Second contact",
     "Legitimate delivery-day check-in - possible lack of visibility",
     "Real-time \"out for delivery\" / live rider location", "Product"),
    ("S3", "After OG EDD", "Postponed", "No Second contact",
     "Order was already late, but agent was able to reassure in one contact",
     "Good outcome for CX - to be improved upon by taking learnings", "CX / Logistics"),
    ("S0", "No EDD data", "N/A", "No Second contact",
     "No EDD revision history exists for this order - nothing to measure a promise against",
     "Investigate why these orders have zero EDD revision rows (dispatch/ERP integration gap)",
     "Product / Data"),
    ("M1", "Before OG EDD", "No Change", "On or Before Revised EDD",
     "Reassurance didn't stick the first time",
     "L1 script: give a specific date + next steps, not generic reassurance", "CX"),
    ("M2", "Before OG EDD", "No Change", "After Revised EDD",
     "Real delay happened, was never caught/communicated between contacts",
     "Proactive alert the moment a delay looks likely", "Product / comms"),
    ("M3", "Before OG EDD", "Postponed", "On or Before Revised EDD",
     "Customer needed a second nudge even though the system had already adjusted",
     "Confirmation SMS/WhatsApp when EDD is revised, so they don't need to ask again",
     "Product / comms"),
    ("M4", "Before OG EDD", "Postponed", "After Revised EDD",
     "The revised promise was missed too",
     "Logistics RCA on repeat-delay orders", "Logistics"),
    ("M5", "On OG EDD", "No Change", "On or Before Revised EDD",
     "Same as M1, triggered on delivery day instead", "Same fix as M1", "CX"),
    ("M6", "On OG EDD", "No Change", "After Revised EDD",
     "Same as M2, triggered on delivery day instead", "Same fix as M2", "Product / comms"),
    ("M7", "On OG EDD", "Postponed", "On or Before Revised EDD",
     "Same as M3, triggered on delivery day instead", "Same fix as M3", "Product / comms"),
    ("M8", "On OG EDD", "Postponed", "After Revised EDD",
     "Same as M4, triggered on delivery day instead", "Same fix as M4", "Logistics"),
    ("M9", "After OG EDD", "No Change (as of 1st contact)", "After Revised EDD",
     "Order was late and the system never once updated the promise across two touches",
     "Highest-priority logistics fix - total silence on a known-late order", "Logistics"),
    ("M10", "After OG EDD", "Postponed", "On or Before Revised EDD",
     "Customer needed extra reassurance even though the (revised) promise held",
     "L1 empowerment: firm new-date commitment + confirmation SMS to close the loop", "CX"),
    ("M11", "After OG EDD", "Postponed", "After Revised EDD",
     "Order was already broken, got a new promise, and broke that promise too",
     "Containment SLA at Logistics - the biggest single lever in this dataset", "Logistics"),
    ("M12", "After OG EDD", "No Change (as of 1st) -> Revised before 2nd", "On or Before Revised EDD",
     "Order was silently overdue at first contact, but caught up with a revised (and honored) "
     "promise before the customer followed up again",
     "Same underlying gap as M9 (delay wasn't flagged proactively) but it self-corrected in "
     "time - tighten the trigger so this is proactive, not lucky timing", "Product / Logistics"),
    ("M0", "No EDD data", "N/A", "N/A",
     "Same data gap as S0, compounded by a repeat contact on an order with no trackable promise",
     "Same fix as S0, prioritized given the repeat-contact signal", "Product / Data"),
    ("MX", "Varies", "Pulled Early (order due sooner than originally promised)", "N/A",
     "Delivery moved up after the first contact - a good-news case, not a failure",
     "None required; optionally notify the customer the order is arriving early", "Logistics"),
]

FIRST_CONTACT_FILL = {"Before OG EDD": GREEN, "On OG EDD": YELLOW, "After OG EDD": ORANGE,
                       "No EDD data": GREY, "Varies": GREY}
EDD_CHANGE_FILL_PREFIX = {"No Change": YELLOW, "Postponed": ORANGE, "N/A": GREY, "Pulled": GREEN}
SECOND_CONTACT_FILL = {"On or Before Revised EDD": GREEN, "After Revised EDD": ORANGE,
                        "No Second contact": GREY, "N/A": GREY}


def fill_for(col, value):
    if col == "first":
        return FIRST_CONTACT_FILL.get(value, GREY)
    if col == "edd":
        for prefix, color in EDD_CHANGE_FILL_PREFIX.items():
            if value.startswith(prefix):
                return color
        return GREY
    if col == "second":
        return SECOND_CONTACT_FILL.get(value, GREY)


def main():
    df = pd.read_parquet("/tmp/edd_with_cohort_v3_fixed.parquet")
    full_cohort = df["cohort"].combine_first(df["scohort"])
    counts = full_cohort.value_counts()
    total = len(df)
    assert total == 119861, f"expected 119861 rows, got {total}"
    assert counts.sum() == total, f"cohort counts sum to {counts.sum()}, expected {total}"

    out_path = PIPELINE_ROOT / "output" / "PRX_Cohort_Definitions_Updated.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        headers = ["#", "First contact", "EDD change", "Second+ contact", "Count",
                   "Overall share", "Root cause", "Fix", "Owner"]
        wb = writer.book
        ws = wb.create_sheet("Cohort Definitions")
        for leftover in [s for s in wb.sheetnames if s != "Cohort Definitions"]:
            del wb[leftover]

        ws.merge_cells("A1:I1")
        ws["A1"] = "PRX CPO Reduction | 'delayed_or_likely_to_be_delayed' & 'shipment_status_query' (updated)"
        ws["A1"].font = Font(bold=True, size=13, color="632423")
        ws["A2"] = "Cohorts Mx - multi-contact orders (contact 1 vs contact 2 only) | Cohorts Sx - single-contact orders"
        ws["A2"].font = Font(italic=True, size=9)

        header_row = 4
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row_i = header_row + 1
        for code, first, edd, second, cause, fix, owner in ROWS:
            count = int(counts.get(code, 0))
            share = count / total
            vals = [code, first, edd, second, count, share, cause, fix, owner]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=row_i, column=c, value=v)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if c == 2:
                    cell.fill = PatternFill("solid", fgColor=fill_for("first", first))
                elif c == 3:
                    cell.fill = PatternFill("solid", fgColor=fill_for("edd", edd))
                elif c == 4:
                    cell.fill = PatternFill("solid", fgColor=fill_for("second", second))
                elif c == 6:
                    cell.number_format = "0.00%"
            row_i += 1

        total_row = row_i
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=int(counts.sum())).font = Font(bold=True)
        tc = ws.cell(row=total_row, column=6, value=1.0)
        tc.number_format = "0.00%"
        tc.font = Font(bold=True)

        widths = [5, 14, 26, 22, 8, 12, 42, 42, 16]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{header_row + 1}"

    print(f"[16] wrote {out_path}")
    print(counts.reindex([r[0] for r in ROWS]))
    print("TOTAL:", int(counts.sum()), "/", total)


if __name__ == "__main__":
    main()
