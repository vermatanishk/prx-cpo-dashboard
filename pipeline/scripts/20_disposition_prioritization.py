"""Stage 20: Disposition Prioritization -- ranks 5 target sub-dispositions (2 Delivery Issues +
3 Return & Pickup) against the full universe of sub-dispositions on two business-facing metrics:

  Metric A - "Interactions per issue" (touches-to-resolution): from resolution_by_issue.parquet
    (Stage 09's 48h-silence-gap resolution definition), flattened touches_per_resolution per
    sub_disposition -> median, mean, % of resolution-episodes needing 3+ touches.

  Metric B - "Contacts per order" volume share: total raw contact count + % share of all support
    contacts + distinct orders affected, from dedupable_raw.parquet. Reuses the EXACT methodology
    Stage 12 (build_subdisp_before_after) already uses for per-sub-disposition contact share --
    value_counts() of normalized_sub_disposition over dedupable_raw, share = count / total
    dedupable_raw count -- rather than re-deriving a possibly-inconsistent denominator.

Both metrics are computed for EVERY sub_disposition in the data (for context/ranking) but only the
5 targets are the headline output.

Reads:  intermediate/run_2026-05-01_2026-08-22/dedupable_raw.parquet
        intermediate/run_2026-05-01_2026-08-22/resolution_by_issue.parquet
        reference/dispositions_master.csv
Writes: output/PRX_Disposition_Prioritization_2026-05-01_to_2026-08-22.xlsx
        ("Priority Matrix (5 target)" + "All Sub-Dispositions (context)")
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from pipeline_utils import PIPELINE_ROOT

WINDOW_DIR = PIPELINE_ROOT / "intermediate" / "run_2026-05-01_2026-08-22"
OUT_PATH = PIPELINE_ROOT / "output" / "PRX_Disposition_Prioritization_2026-05-01_to_2026-08-22.xlsx"

HEADER_FILL = "632423"

# Confirmed via reference/dispositions_master.csv's sub_disposition_normalized column.
TARGET_SUBDISPS = [
    "delayed_or_likely_to_be_delayed",         # Delivery Issues - Delayed or likely to be delayed
    "shipment_status_query_(no_delay)",        # Delivery Issues - Shipment status query (no delay)
    "spilled_or_broken_or_spoiled_contents",   # Return & Pickup - cold chain item damage
    "different_item_delivered_than_ordered",   # Return & Pickup - wrong item received
    "missing_or_less_quantity_received",       # Return & Pickup - missing qty
]

OUTPUT_COLS = [
    "Disposition", "Sub-Disposition", "Total Contacts", "% Share of All Contacts",
    "Distinct Orders Affected", "Median Touches-to-Resolution", "Mean Touches-to-Resolution",
    "% Resolutions Needing 3+ Touches", "Resolution Episodes (n)",
]


def autosize_and_freeze(writer, sheet_name, df):
    ws = writer.sheets[sheet_name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        width = min(60, max(10, int(df[col].astype(str).map(len).clip(upper=60).max() if len(df) else 10) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width


def style_header(writer, sheet_name, n_cols):
    ws = writer.sheets[sheet_name]
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_disposition_map():
    """sub_disposition_normalized -> (Disposition, Sub-Disposition display label). A handful of
    normalized keys (the 8 "issue_not_listed" catch-alls) map to >1 parent Disposition in the
    master taxonomy -- keep the first occurrence for a stable single label, since this map is for
    display/context only and doesn't affect any of the 5 target rows."""
    dm = pd.read_csv(PIPELINE_ROOT / "reference" / "dispositions_master.csv")
    dm = dm.drop_duplicates(subset="sub_disposition_normalized", keep="first")
    return dm.set_index("sub_disposition_normalized")[["disposition", "sub_disposition"]]


def build_metrics(dedupable: pd.DataFrame, resolution: pd.DataFrame, disp_map: pd.DataFrame) -> pd.DataFrame:
    total_contacts = len(dedupable)
    contact_counts = dedupable["normalized_sub_disposition"].value_counts()
    distinct_orders = dedupable.groupby("normalized_sub_disposition")["final_order_id"].nunique()

    exploded = resolution[["sub_disposition", "touches_per_resolution"]].explode("touches_per_resolution")
    exploded["touches_per_resolution"] = exploded["touches_per_resolution"].astype(int)
    grp = exploded.groupby("sub_disposition")["touches_per_resolution"]
    median_touches = grp.median()
    mean_touches = grp.mean()
    pct_3plus = grp.apply(lambda s: (s >= 3).mean())
    n_episodes = grp.size()

    all_subdisps = sorted(set(contact_counts.index) | set(median_touches.index))
    rows = []
    for sub in all_subdisps:
        if sub in disp_map.index:
            disp_label = disp_map.loc[sub, "disposition"]
            subdisp_label = disp_map.loc[sub, "sub_disposition"]
        else:
            disp_label = "(unmapped in dispositions_master.csv)"
            subdisp_label = sub
        rows.append({
            "Disposition": disp_label,
            "Sub-Disposition": subdisp_label,
            "sub_disposition_normalized": sub,
            "Total Contacts": int(contact_counts.get(sub, 0)),
            "% Share of All Contacts": contact_counts.get(sub, 0) / total_contacts,
            "Distinct Orders Affected": int(distinct_orders.get(sub, 0)),
            "Median Touches-to-Resolution": median_touches.get(sub, pd.NA),
            "Mean Touches-to-Resolution": mean_touches.get(sub, pd.NA),
            "% Resolutions Needing 3+ Touches": pct_3plus.get(sub, pd.NA),
            "Resolution Episodes (n)": int(n_episodes.get(sub, 0)),
        })
    out = pd.DataFrame(rows).sort_values("Total Contacts", ascending=False).reset_index(drop=True)
    out.insert(0, "Rank by Total Contacts", range(1, len(out) + 1))
    return out


def format_numeric_sheet(writer, sheet_name, df):
    ws = writer.sheets[sheet_name]
    col_idx = {c: i for i, c in enumerate(df.columns, start=1)}
    pct_cols = [c for c in ["% Share of All Contacts", "% Resolutions Needing 3+ Touches"] if c in col_idx]
    mean_cols = [c for c in ["Mean Touches-to-Resolution"] if c in col_idx]
    for r in range(2, len(df) + 2):
        for c in pct_cols:
            ws.cell(row=r, column=col_idx[c]).number_format = "0.0%"
        for c in mean_cols:
            ws.cell(row=r, column=col_idx[c]).number_format = "0.00"


def main():
    dedupable = pd.read_parquet(WINDOW_DIR / "dedupable_raw.parquet")
    resolution = pd.read_parquet(WINDOW_DIR / "resolution_by_issue.parquet")
    print(f"[20] dedupable_raw: {len(dedupable)} contacts, window "
          f"{dedupable['created_at'].min()} -> {dedupable['created_at'].max()}")
    print(f"[20] resolution_by_issue: {len(resolution)} (order, sub-disposition, month) keys")

    disp_map = build_disposition_map()
    all_df = build_metrics(dedupable, resolution, disp_map)

    target_df = all_df[all_df["sub_disposition_normalized"].isin(TARGET_SUBDISPS)].copy()
    target_df = target_df.sort_values("Total Contacts", ascending=False).reset_index(drop=True)

    missing = set(TARGET_SUBDISPS) - set(target_df["sub_disposition_normalized"])
    if missing:
        print(f"[20] WARNING: target sub-dispositions not found in data: {missing}")

    priority_out = target_df[OUTPUT_COLS].copy()
    all_out = all_df[["Rank by Total Contacts"] + OUTPUT_COLS].copy()

    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        priority_out.to_excel(writer, sheet_name="Priority Matrix (5 target)", index=False)
        autosize_and_freeze(writer, "Priority Matrix (5 target)", priority_out)
        style_header(writer, "Priority Matrix (5 target)", len(priority_out.columns))
        format_numeric_sheet(writer, "Priority Matrix (5 target)", priority_out)

        all_out.to_excel(writer, sheet_name="All Sub-Dispositions (context)", index=False)
        autosize_and_freeze(writer, "All Sub-Dispositions (context)", all_out)
        style_header(writer, "All Sub-Dispositions (context)", len(all_out.columns))
        format_numeric_sheet(writer, "All Sub-Dispositions (context)", all_out)

    print(f"[20] wrote {OUT_PATH}")
    print(f"[20] {len(all_df)} distinct sub-dispositions in universe")
    print("[20] Priority Matrix (5 target):")
    print(priority_out.to_string(index=False))
    print("[20] Rank of each target sub-disposition by Total Contacts (of {} total):".format(len(all_df)))
    for _, row in target_df.iterrows():
        print(f"    {row['sub_disposition_normalized']}: rank #{row['Rank by Total Contacts']} "
              f"of {len(all_df)}, {row['Resolution Episodes (n)']} resolution episodes")


if __name__ == "__main__":
    main()
