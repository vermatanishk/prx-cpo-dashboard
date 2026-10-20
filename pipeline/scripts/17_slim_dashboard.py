"""Stage 17: Slim Dashboard -- rolling Summary, reconciled Contacts-to-Resolution (with inline
Overall FCR / BFCR / HFCR), and a simplified Sub-Disposition tab.

Format locked to the EIR's own manually-edited reference (screenshots, 2026-09-29): Summary as one
rolling table with months in REVERSE chronological order (latest first) + Total; Contacts-to-
Resolution's CUT 1/CUT 2 histograms with a small FCR/BFCR/HFCR panel computed directly from each
cut's own "1 touch" bucket (no separate BCR/HFCR/FCR tab -- that was tried twice and replaced);
Sub-Disposition rebuilt to a 6-column volume+share+IPI spec. The 0-touches-bucket reconciliation to
Total Orders (added for a prior request) is DROPPED here since the reference sheet doesn't have it
-- flagged in Methodology in case it's still wanted back.

IPO / IPI / CPO decomposition (unchanged since the first build of this stage):
  IPO (Issues per Order)       = Unique issues (post-dedup) / Total Orders
                                  -- "how often something goes wrong" (product/ops lever). This IS
                                     the old "CPO Overall" definition -- labelled as such in Summary.
  IPI (Interactions per Issue) = Total raw interactions (pre-dedup) / Unique issues
                                  -- "how badly we handle it once it happens" (resolution /
                                     conversation-quality lever)
  CPO = IPO x IPI               = Total raw interactions (pre-dedup) / Total Orders
"Issue" = one (Final Order ID x normalized Sub-disposition) dedup group -- Stage 08's existing key,
not redefined here.

Contacts-to-Resolution histogram bucketing bug (found and fixed 2026-09, still in effect): Stage
12's own `_resolution_cut` buckets by RESOLUTION EPISODE, not by key -- a key that recurred after a
>48h gap contributed multiple rows, so its histogram's bucket-sum silently exceeded its own "Total
keys" stat. This dashboard uses a corrected per-key aggregation (each key's FIRST resolution
episode's touch count, exactly one row per key) via `build_resolution_cuts_by_key` below -- NOT
Stage 12's `build_contacts_to_resolution`. Stage 12's own output was not touched.

Overall FCR / BFCR / HFCR (per cut, matching the EIR's own worked formula from the reference
sheet): computed directly off that cut's "1 touch" bucket row, all three sharing the SAME
denominator (that cut's own Total keys) -- NOT the earlier (rejected) version that used different
denominators per metric and a separate resolution_by_issue.parquet population.
  Overall FCR = (keys resolved in exactly 1 touch) / (Total keys)
  BFCR        = (keys resolved in exactly 1 touch, bot-only)  / (Total keys)
  HFCR        = (keys resolved in exactly 1 touch, human-touched) / (Total keys)
  BFCR + HFCR = Overall FCR, by construction (a clean partition of the "1 touch" bucket).

Sub-Disposition tab: Sub-Disposition | Total Contacts with this sub-disp (before dedup total) |
before dedup share | Total Orders with this sub-disp (after dedup total) | after dedup share | IPI
(Contacts/Orders).

Reads:  intermediate/.../dedupable_raw.parquet            (Before Dedup)
        intermediate/.../deduped_primary_contacts.parquet  (After Dedup = "Issues")
        intermediate/.../resolution_by_order.parquet       (Order-level 48h resolution, Stage 09)
        intermediate/.../fallback_orders.parquet           (order -> user_id map, for user metrics)
        raw/orders/orders_denominator.csv
Writes: output/PRX_Slim_Dashboard_<start>_to_<end>_<version_tag>.xlsx
"""
import importlib.util
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from pipeline_utils import load_config, intermediate_dir, PIPELINE_ROOT

# Reuse Stage 09's `contacts_to_resolution` (the atomic per-key touches-to-resolution function,
# imported here via Stage 12, which already imports it) -- the correct, already-decided logic.
_spec = importlib.util.spec_from_file_location(
    "cpo_disposition_views", Path(__file__).parent / "12_cpo_disposition_views.py"
)
_cpo_views = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_cpo_views)
contacts_to_resolution = _cpo_views.contacts_to_resolution

CHANNEL_ORDER = [
    "Ringg", "Agent IB", "Agent OB WH (connected)", "LimeChat Bot",
    "LimeChat Bot + Agent", "Email", "Agent OB WH (completed)", "Agent OB WH (other)",
]

HEADER_FILL = PatternFill("solid", fgColor="632423")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FONT = Font(bold=True, size=12)


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def autosize(ws, df, start_col=1):
    for i, col in enumerate(df.columns, start=start_col):
        width = min(60, max(10, int(df[col].astype(str).map(len).clip(upper=60).max() if len(df) else 10) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width


def write_table(writer, sheet_name, df, title=None, startrow=0, startcol=0, freeze=True, pct_cols=None):
    if title:
        pd.DataFrame({title: []}).to_excel(writer, sheet_name=sheet_name, startrow=startrow, startcol=startcol, index=False)
        ws = writer.sheets[sheet_name]
        ws.cell(row=startrow + 1, column=startcol + 1).font = SECTION_FONT
        startrow += 2
    df.to_excel(writer, sheet_name=sheet_name, startrow=startrow, startcol=startcol, index=False)
    ws = writer.sheets[sheet_name]
    style_header_row(ws, startrow + 1, len(df.columns), start_col=startcol + 1)
    if pct_cols:
        for col_name in pct_cols:
            if col_name in df.columns:
                col_idx = startcol + list(df.columns).index(col_name) + 1
                for r in range(startrow + 2, startrow + 2 + len(df)):
                    ws.cell(row=r, column=col_idx).number_format = "0.00%"
    if freeze and startrow == 0 and startcol == 0:
        ws.freeze_panes = "A2"
    autosize(ws, df, start_col=startcol + 1)
    return startrow + len(df) + 3


def month_label(m):
    return pd.Period(m, freq="M").strftime("%b %Y")


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

def load_data(cfg):
    out_dir = intermediate_dir(cfg)
    dedupable = pd.read_parquet(out_dir / "dedupable_raw.parquet")
    primaries = pd.read_parquet(out_dir / "deduped_primary_contacts.parquet")
    resolution_order = pd.read_parquet(out_dir / "resolution_by_order.parquet")
    fallback_orders = pd.read_parquet(out_dir / "fallback_orders.parquet")
    orders_denom = pd.read_csv(PIPELINE_ROOT / "raw" / "orders" / "orders_denominator.csv",
                                parse_dates=["order_date"])
    dedupable["created_at"] = pd.to_datetime(dedupable["created_at"], errors="coerce")
    dedupable["month"] = dedupable["created_at"].dt.to_period("M").astype(str)
    primaries["created_at"] = pd.to_datetime(primaries["created_at"], errors="coerce")
    primaries["month"] = primaries["created_at"].dt.to_period("M").astype(str)
    fallback_orders["month"] = fallback_orders["created_at"].dt.to_period("M").astype(str)
    return dedupable, primaries, resolution_order, fallback_orders, orders_denom


# ---------------------------------------------------------------------------
# Summary tab -- ONE rolling Metric x Month table, months latest-first
# ---------------------------------------------------------------------------

def build_summary_table(dedupable, primaries, resolution_order, fallback_orders, orders_denom):
    orders_denom = orders_denom.copy()
    orders_denom["month"] = orders_denom["order_date"].dt.to_period("M").astype(str)
    orders_by_month = orders_denom.groupby("month")["total_orders_processed"].sum()
    months_asc = sorted(orders_by_month.index)
    months = list(reversed(months_asc))  # latest first, per reference format
    cols = months + ["Total"]

    orders_w_contact = resolution_order.groupby("month").size()
    orders_w_contact_actionable = primaries[primaries["is_actionable"]].groupby("month")["final_order_id"].nunique()

    order_to_user = fallback_orders.drop_duplicates("order_id").set_index("order_id")["user_id"]
    total_unique_users = fallback_orders.groupby("month")["user_id"].nunique()

    orders_with_contact_by_month = resolution_order.groupby("month")["final_order_id"].apply(set)
    users_with_contact = {}
    for m in months_asc:
        order_ids = orders_with_contact_by_month.get(m, set())
        users_with_contact[m] = pd.Series(list(order_ids)).map(order_to_user).dropna().nunique() if order_ids else 0

    total_contacts_raw = dedupable.groupby("month").size()
    total_contacts_after_dedup = primaries.groupby("month").size()

    rows = []

    def add_row(metric, series_or_dict, is_pct=False, numerator=None, denominator=None):
        row = {"Metric": metric}
        for m in months:
            row[m] = series_or_dict.get(m, 0) if isinstance(series_or_dict, dict) else series_or_dict.get(m, 0)
        if is_pct:
            for m in months:
                num = numerator.get(m, 0)
                den = denominator.get(m, 0)
                row[m] = (num / den) if den else 0
            total_num = sum(numerator.get(m, 0) for m in months_asc)
            total_den = sum(denominator.get(m, 0) for m in months_asc)
            row["Total"] = (total_num / total_den) if total_den else 0
        else:
            row["Total"] = sum(row[m] for m in months)
        rows.append(row)
        return row

    add_row("Total Orders", orders_by_month)
    add_row("Total Orders w Support Contact", orders_w_contact)
    add_row("Total Orders w Support Contact - Actionable", orders_w_contact_actionable)
    add_row("% Support Driven Orders", {}, is_pct=True, numerator=orders_w_contact, denominator=orders_by_month)
    add_row("Total Unique Users (all orders)", total_unique_users)
    add_row("Total Unique Users w/ Support Contact", pd.Series(users_with_contact))
    add_row("% Users Who Contact for Support at Least Once", {}, is_pct=True,
            numerator=pd.Series(users_with_contact), denominator=total_unique_users)
    add_row("Total Contacts Made (raw, support-only, sales stripped)", total_contacts_raw)
    add_row("Total Contacts After De-dup (new logic)", total_contacts_after_dedup)

    rows.append({"Metric": "", **{m: None for m in cols}})
    rows.append({"Metric": "IPO (Unique Issues / Total Orders) [old definition of CPO]", **{m: None for m in cols}})
    add_row_ipo = {"Metric": "IPO (Unique Issues / Total Orders) [old definition of CPO]"}
    add_row_ipi = {"Metric": "IPI (Total Interactions / Unique Issues)"}
    add_row_cpo = {"Metric": "CPO = IPO x IPI"}
    rows.pop()  # remove placeholder header, real row follows below
    for m in months:
        orders_m = orders_by_month.get(m, 0)
        issues_m = total_contacts_after_dedup.get(m, 0)
        interactions_m = total_contacts_raw.get(m, 0)
        add_row_ipo[m] = (issues_m / orders_m) if orders_m else 0
        add_row_ipi[m] = (interactions_m / issues_m) if issues_m else 0
        add_row_cpo[m] = (interactions_m / orders_m) if orders_m else 0
    total_orders_all = orders_by_month.sum()
    total_issues_all = total_contacts_after_dedup.sum()
    total_interactions_all = total_contacts_raw.sum()
    add_row_ipo["Total"] = (total_issues_all / total_orders_all) if total_orders_all else 0
    add_row_ipi["Total"] = (total_interactions_all / total_issues_all) if total_issues_all else 0
    add_row_cpo["Total"] = (total_interactions_all / total_orders_all) if total_orders_all else 0
    rows.extend([add_row_ipo, add_row_ipi, add_row_cpo])

    rows.append({"Metric": "", **{m: None for m in cols}})

    before_pivot = dedupable.pivot_table(index="sub_channel", columns="month", aggfunc="size", fill_value=0)
    after_pivot = primaries.pivot_table(index="sub_channel", columns="month", aggfunc="size", fill_value=0)

    rows.append({"Metric": "Contact Volume by Channel -- Before Dedup (raw, support-only, sales stripped)",
                  **{m: None for m in cols}})
    for ch in CHANNEL_ORDER:
        row = {"Metric": ch}
        for m in months:
            row[m] = int(before_pivot.loc[ch, m]) if ch in before_pivot.index and m in before_pivot.columns else 0
        row["Total"] = sum(row[m] for m in months)
        rows.append(row)

    rows.append({"Metric": "", **{m: None for m in cols}})
    rows.append({"Metric": "Contact Volume by Channel -- After Dedup (new logic: Order ID x Sub-disposition, first-Actionable)",
                  **{m: None for m in cols}})
    for ch in CHANNEL_ORDER:
        row = {"Metric": ch}
        for m in months:
            row[m] = int(after_pivot.loc[ch, m]) if ch in after_pivot.index and m in after_pivot.columns else 0
        row["Total"] = sum(row[m] for m in months)
        rows.append(row)

    df = pd.DataFrame(rows)
    rename = {m: month_label(m) for m in months}
    df = df.rename(columns=rename)
    df = df[["Metric"] + [month_label(m) for m in months] + ["Total"]]
    return df, months_asc


# ---------------------------------------------------------------------------
# Contacts-to-Resolution -- per-key-corrected histogram + inline FCR/BFCR/HFCR panel
# ---------------------------------------------------------------------------

def _resolution_cut_by_key(dedupable, window_hours, group_cols, order_id_label,
                            bot_channels=("Ringg", "LimeChat Bot")):
    rows = []
    for _, g in dedupable.groupby(group_cols):
        touches = contacts_to_resolution(g, window_hours)
        bot_only = g["sub_channel"].isin(list(bot_channels)).all()
        rows.append({"touches": touches[0], "bot_contained_whole_journey": bot_only})
    events = pd.DataFrame(rows)

    bucketed = events["touches"].clip(upper=5).astype(str).where(events["touches"] < 5, "5+")
    hist = events.groupby(bucketed).agg(
        total=("touches", "size"),
        bot_resolved=("bot_contained_whole_journey", "sum"),
    )
    hist["human_resolved"] = hist["total"] - hist["bot_resolved"]
    hist["pct_of_total"] = hist["total"] / hist["total"].sum()
    order = [str(i) for i in range(1, 5)] + ["5+"]
    hist = hist.reindex(order).fillna(0)
    hist["cumulative_pct"] = hist["pct_of_total"].cumsum()
    hist["bot_containment"] = hist["bot_resolved"] / hist["total"].replace(0, pd.NA)
    hist = hist.reset_index().rename(columns={
        "touches": "# contacts", "total": order_id_label, "pct_of_total": "% of Total",
        "cumulative_pct": "Cumulative %", "bot_resolved": "Bot Resolved",
        "human_resolved": "Human Resolved", "bot_containment": "Bot containment",
    })

    total_keys = len(events)
    stats = {
        "Total keys": total_keys,
        "Mean": round(events["touches"].mean(), 2),
        "Median (P50)": events["touches"].median(),
        "P75": events["touches"].quantile(0.75),
        "P90": events["touches"].quantile(0.90),
        "P95": events["touches"].quantile(0.95),
        "P99": events["touches"].quantile(0.99),
        "P99.9": events["touches"].quantile(0.999),
        "Max (tail)": events["touches"].max(),
    }

    # FCR / BFCR / HFCR, computed directly off the "1 touch" bucket, all three sharing the SAME
    # denominator (this cut's own Total keys) -- matches the EIR's own worked formula exactly.
    row1 = hist[hist["# contacts"] == "1"].iloc[0]
    fcr = {
        "Overall FCR": row1["% of Total"],
        "BFCR": (row1["Bot Resolved"] / total_keys) if total_keys else 0,
        "HFCR": (row1["Human Resolved"] / total_keys) if total_keys else 0,
    }
    return hist, stats, fcr


def build_resolution_cuts_by_key(dedupable, window_hours):
    cut1_hist, cut1_stats, cut1_fcr = _resolution_cut_by_key(
        dedupable, window_hours, ["final_order_id", "month"], "Total OrderIDs"
    )
    cut2_pop = dedupable[dedupable["is_actionable"]]
    cut2_hist, cut2_stats, cut2_fcr = _resolution_cut_by_key(
        cut2_pop, window_hours, ["final_order_id", "normalized_sub_disposition", "month"], "Total Keys"
    )
    return [
        ("CUT 1 - Order ID level (all support-only contacts per order)", cut1_hist, cut1_stats, cut1_fcr),
        ("CUT 2 - Order ID x Sub-disposition level, Actionable only", cut2_hist, cut2_stats, cut2_fcr),
    ]


def write_resolution_sheet_with_fcr(writer, sheet_name, cuts):
    startrow = 0
    for label, hist, stats, fcr in cuts:
        pd.DataFrame({label: []}).to_excel(writer, sheet_name=sheet_name, startrow=startrow, index=False)
        writer.sheets[sheet_name].cell(row=startrow + 1, column=1).font = SECTION_FONT
        startrow += 2

        hist_cols = ["# contacts"] + [c for c in hist.columns if c != "# contacts"]
        hist = hist[hist_cols]
        hist.to_excel(writer, sheet_name=sheet_name, startrow=startrow, index=False)
        ws = writer.sheets[sheet_name]
        style_header_row(ws, startrow + 1, len(hist.columns))
        for col_name in ["% of Total", "Cumulative %", "Bot containment"]:
            col_idx = list(hist.columns).index(col_name) + 1
            for r in range(startrow + 2, startrow + 2 + len(hist)):
                ws.cell(row=r, column=col_idx).number_format = "0.00%"

        stats_df = pd.DataFrame({"Stat": list(stats.keys()), "Value": list(stats.values())})
        write_table(writer, sheet_name, stats_df, startrow=startrow, startcol=8, freeze=False)

        fcr_df = pd.DataFrame({"Metric": list(fcr.keys()), "Value": list(fcr.values())})
        write_table(writer, sheet_name, fcr_df, startrow=startrow, startcol=11, freeze=False, pct_cols=["Value"])

        startrow += len(hist) + 4


# ---------------------------------------------------------------------------
# Sub-Disposition tab -- volume + share + IPI
# ---------------------------------------------------------------------------

def build_sub_disposition_table(dedupable, primaries):
    before = dedupable.groupby("normalized_sub_disposition").size().rename("Total Contacts with this sub-disp (before dedup total)")
    after = primaries.groupby("normalized_sub_disposition").size().rename("Total Orders with this sub-disp (after dedup total)")
    out = pd.concat([before, after], axis=1).fillna(0)
    before_col = "Total Contacts with this sub-disp (before dedup total)"
    after_col = "Total Orders with this sub-disp (after dedup total)"
    out[before_col] = out[before_col].astype(int)
    out[after_col] = out[after_col].astype(int)
    out["before dedup share"] = out[before_col] / out[before_col].sum()
    out["after dedup share"] = out[after_col] / out[after_col].sum()
    out["IPI (Contacts/Orders)"] = (out[before_col] / out[after_col].replace(0, pd.NA)).fillna(0)
    out.index.name = "Sub-Disposition"
    out = out.reset_index().sort_values(before_col, ascending=False)
    return out[["Sub-Disposition", before_col, "before dedup share", after_col, "after dedup share",
                "IPI (Contacts/Orders)"]]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    cfg = load_config()
    dedupable, primaries, resolution_order, fallback_orders, orders_denom = load_data(cfg)

    total_orders = orders_denom["total_orders_processed"].sum()
    print(f"[17] Total Orders in period: {total_orders}")
    print(f"[17] Before Dedup (dedupable_raw): {len(dedupable)} rows")
    print(f"[17] After Dedup / Unique Issues (deduped_primary_contacts): {len(primaries)} rows")

    summary, months = build_summary_table(dedupable, primaries, resolution_order, fallback_orders, orders_denom)

    resolution_cuts = build_resolution_cuts_by_key(dedupable, cfg["resolution"]["window_hours"])
    for label, hist, stats, fcr in resolution_cuts:
        bucket_sum = hist[[c for c in hist.columns if "Total" in c][0]].sum()
        print(f"[17] {label}: bucket sum={int(bucket_sum)}, Total keys stat={stats['Total keys']} "
              f"(should match) | FCR={fcr['Overall FCR']:.2%} BFCR={fcr['BFCR']:.2%} HFCR={fcr['HFCR']:.2%}")

    subdisp = build_sub_disposition_table(dedupable, primaries)

    start, end = cfg["run"]["start_date"], cfg["run"]["end_date"]
    version_tag = cfg["run"].get("version_tag", "")
    out_path = PIPELINE_ROOT / "output" / f"PRX_Slim_Dashboard_{start}_to_{end}_{version_tag}.xlsx"

    methodology = pd.DataFrame({"Definition": [
        "IPO (Issues per Order) = Unique Issues (post-dedup, Order ID x Sub-disposition) / Total Orders "
        "in the period. Answers: how often does something go wrong. This IS the old 'CPO Overall' "
        "definition -- labelled as such directly in the Summary tab's row, rather than kept as a "
        "separate legacy row.",
        "IPI (Interactions per Issue) = Total raw interactions (pre-dedup) / Unique Issues. Answers: "
        "how badly do we handle it once it happens. A resolution/conversation-quality lever.",
        "CPO = IPO x IPI = Total raw interactions (pre-dedup) / Total Orders. Identity, not "
        "independently re-derived -- IPO x IPI always reconciles to CPO by construction.",
        "'Issue' = one (Final Order ID x normalized Sub-disposition) dedup group -- Stage 08's "
        "existing dedup key. Not redefined here.",
        "Summary tab months are ordered LATEST FIRST (matching the reference format), with a Total "
        "column last. 'Total Orders w Support Contact' is sourced from resolution_by_order.parquet "
        "(one row per (Order ID, Month) key), so it ties out exactly to Contacts-to-Resolution "
        "CUT 1's Total keys stat by construction.",
        "Contacts-to-Resolution histogram bucketing BUG (found and fixed): Stage 12's own "
        "_resolution_cut bucketed by RESOLUTION EPISODE, not by key -- a key that recurred after a "
        ">48h gap contributed multiple rows, so the histogram's bucket-sum silently exceeded its own "
        "'Total keys' stat. This dashboard uses a corrected per-key aggregation instead (each key's "
        "FIRST resolution episode's touch count, exactly one row per key). Stage 12's own output was "
        "NOT touched -- flag separately if PRX_CPO_Analysis_v4.xlsx needs the same fix.",
        "DROPPED vs. an earlier version of this dashboard: the '0 touches (no support contact)' "
        "bucket that reconciled CUT 1's grand total to Total Orders (897k-style). The EIR's own "
        "reference sheet for this rebuild doesn't include it -- removed to match. Flag if it's still "
        "wanted back.",
        "Overall FCR / BFCR / HFCR (per cut) -- REPLACES the earlier separate 'BCR-HFCR-FCR' tab "
        "(tried twice, both rejected). Computed directly off that cut's own '1 touch' bucket, all "
        "three sharing the SAME denominator (that cut's Total keys), matching the EIR's own worked "
        "formula from the reference sheet: Overall FCR = (keys resolved in 1 touch) / (Total keys); "
        "BFCR = (keys resolved in 1 touch, bot-only) / (Total keys); HFCR = (keys resolved in 1 "
        "touch, human-touched) / (Total keys). BFCR + HFCR = Overall FCR by construction.",
        "Sub-Disposition tab: 'Total Contacts...(before dedup total)' and 'Total Orders...(after "
        "dedup total)' are literal row counts of dedupable_raw / deduped_primary_contacts for that "
        "sub-disposition, each with its own share-of-total column. IPI = before / after.",
        f"Run range: {start} to {end} (exclusive), version_tag={version_tag}. LimeChat, Orders and "
        "Calls all cover the full range. Zoho coverage remains partial (see "
        "PROJECT_METHODOLOGY_AND_LEARNINGS.md). The latest month is right-censored: orders placed "
        "late in the month have not yet had time to generate contacts, so its IPO/CPO read low -- "
        "compare against the matured-cohort file, not the other months.",
        "Two double-counting fixes in 05_load_calls.py apply from this version on, so numbers are not "
        "directly comparable to v4 and earlier: (1) Agent IB: a call linked to several CRM tasks "
        "('sister tasks') now takes the task carrying a real disposition, latest task_id first -- "
        "previously an arbitrary pick, usually the generic 'Sales Related - Created a New Task' "
        "hand-off marker. (2) Agent Outbound: auto-closed crm_tasks (status != 20) with no "
        "disposition of their own are dropped when a Ringg contact exists on the same order within "
        "15 minutes -- they were back-filled with the Ringg call's own outcome and counted twice.",
    ]})

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        write_table(writer, "Summary", summary)

        write_resolution_sheet_with_fcr(writer, "Contacts-to-Resolution", resolution_cuts)

        write_table(writer, "Sub-Disposition", subdisp,
                    pct_cols=["before dedup share", "after dedup share"])

        write_table(writer, "Methodology", methodology, freeze=False)
        writer.sheets["Methodology"].column_dimensions["A"].width = 120
        for r in range(1, len(methodology) + 2):
            writer.sheets["Methodology"].cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    print(f"[17] wrote {out_path}")


if __name__ == "__main__":
    main()
